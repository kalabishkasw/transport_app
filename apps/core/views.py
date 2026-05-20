"""
веб-інтерфейс диспетчера: дашборд, списки і деталі сутностей.
усі сторінки потребують авторизації (staff_required).
"""

from datetime import timedelta
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.core.paginator import Paginator
from django.db.models import Count, Q, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone

from apps.customers.models import Customer
from apps.fleet.models import Driver, Vehicle
from apps.orders.models import Order, PromoCode, Ticket
from apps.routes.models import Route, Trip

from .decorators import staff_required
from .models import AuditLog


# rate-limit декоратор. no-op якщо django-ratelimit не встановлено.
try:
    from django_ratelimit.decorators import ratelimit
except ImportError:
    def ratelimit(*args, **kwargs):
        def _wrap(view_func):
            return view_func
        return _wrap


PAGE_SIZE = 20



# дашборд


@staff_required
def dashboard(request):
    now = timezone.now()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    week_ahead = now + timedelta(days=7)

    next_month_start = (month_start + timedelta(days=32)).replace(day=1)

    trips_this_month = Trip.objects.filter(
        departure_at__gte=month_start,
        departure_at__lt=next_month_start,
    ).count()
    trips_active_today = Trip.objects.filter(
        departure_at__gte=today_start,
        departure_at__lt=today_start + timedelta(days=1),
    ).count()
    # унікальні авто, задіяні у сьогоднішніх рейсах. може бути менше за trips_active_today,
    # бо одне авто може робити 2-3 рейси на день. використовується для розрахунку
    # відсотка завантаженості автопарку (а не трип-навантаження).
    vehicles_in_use_today = (
        Trip.objects
        .filter(
            departure_at__gte=today_start,
            departure_at__lt=today_start + timedelta(days=1),
        )
        .values('vehicle_id')
        .distinct()
        .count()
    )

    # квитки на рейси цього місяця (більш предметна метрика для бізнесу)
    tickets_this_month = Ticket.objects.filter(
        order__trip__departure_at__gte=month_start,
        order__trip__departure_at__lt=next_month_start,
        status__in=['booked', 'paid', 'used'],
    ).count()

    # виручка з рейсів цього місяця
    revenue_this_month = Order.objects.filter(
        trip__departure_at__gte=month_start,
        trip__departure_at__lt=next_month_start,
        status__in=['paid', 'in_progress', 'completed'],
    ).aggregate(total=Sum('total_price'))['total'] or Decimal('0')

    customers_count = Customer.objects.filter(is_active=True).count()
    vehicles_active = Vehicle.objects.filter(is_active=True).count()
    drivers_available = Driver.objects.filter(is_available=True).count()
    routes_active = Route.objects.filter(is_active=True).count()

    upcoming_trips = (
        Trip.objects
        .filter(departure_at__gte=now, departure_at__lte=week_ahead)
        .select_related('route', 'vehicle', 'main_driver__user')
        .annotate(tickets_count=Count(
            'orders__tickets',
            filter=Q(orders__tickets__status__in=['booked', 'paid', 'used']),
        ))
        .order_by('departure_at')[:8]
    )

    recent_orders = (
        Order.objects
        .select_related('trip__route', 'customer')
        .annotate(tickets_total=Count('tickets'))
        .order_by('-created_at')[:8]
    )

    # дані для графіків

    # виручка по днях: останні 30 днів (актуал) + наступні 60 днів (прогноз).
    # тут важливво: і фактична, і прогнозна частини рахуються за однаковою метрикою -
    # за датою відправлення рейсу. це гарантує плавний перехід через "Сьогодні"
    # без штучного провалу.
    period_start = (now - timedelta(days=29)).replace(hour=0, minute=0, second=0, microsecond=0)
    forecast_end = period_start + timedelta(days=90)

    revenue_by_day = (
        Order.objects
        .filter(
            trip__departure_at__gte=period_start,
            trip__departure_at__lt=forecast_end,
        )
        .exclude(status__in=['cancelled', 'refunded'])
        .annotate(day=TruncDate('trip__departure_at'))
        .values('day')
        .annotate(total=Sum('total_price'))
        .order_by('day')
    )
    revenue_map = {r['day'].isoformat(): float(r['total'] or 0) for r in revenue_by_day}

    # будуємо ряд по 90 днях. точка "Сьогодні" розділяє факт і прогноз
    # тільки візуально (пунктирна лінія у frontend), але метрика одна.
    combined_raw = []
    labels = []
    today_iso = now.date().isoformat()
    today_index = None
    for i in range(90):
        d = (period_start + timedelta(days=i)).date()
        iso = d.isoformat()
        labels.append(d.strftime('%d.%m'))
        if iso == today_iso:
            today_index = i
        combined_raw.append(revenue_map.get(iso, 0))

    def smooth(values, window=7):
        """7-денне центральне ковзне середнє."""
        n = len(values)
        half = window // 2
        result = []
        for i in range(n):
            start = max(0, i - half)
            end = min(n, i + half + 1)
            chunk = [v for v in values[start:end] if v is not None]
            if chunk:
                result.append(round(sum(chunk) / len(chunk), 2))
            else:
                result.append(0)
        return result

    smoothed = smooth(combined_raw)

    chart_revenue = [
        {'label': labels[i], 'value': smoothed[i]}
        for i in range(90)
    ]
    # today_index використовуємо у frontend для відображення вертикальної лінії "сьогодні"
    chart_today_index = today_index if today_index is not None else 30

    # розподіл рейсів за статусом (усі рейси, не лише найближчі).
    # передаємо також `key` (англ. ключ статусу), щоб JS міг призначати
    # кольори за змістом, а не за порядком: cancelled - червоний, completed - зелений тощо.
    status_qs = (
        Trip.objects
        .values('status')
        .annotate(count=Count('id'))
        .order_by()
    )
    status_labels_map = dict(Trip.Status.choices)
    chart_status = [
        {
            'key': s['status'],
            'label': status_labels_map.get(s['status'], s['status']),
            'value': s['count'],
        }
        for s in status_qs
    ]

    # топ-5 найпопулярніших маршрутів (за кількістю проданих квитків).
    # найважчий запит дашборда (joins по 4 таблицях). Кешуємо на 5 хвилин.
    chart_top_routes = cache.get('dashboard_top_routes_v1')
    if chart_top_routes is None:
        top_routes_qs = (
            Route.objects
            .annotate(tickets_count=Count(
                'trips__orders__tickets',
                filter=Q(trips__orders__tickets__status__in=['booked', 'paid', 'used']),
            ))
            .filter(tickets_count__gt=0)
            .order_by('-tickets_count')[:5]
        )
        chart_top_routes = [
            {'label': r.name, 'value': r.tickets_count}
            for r in top_routes_qs
        ]
        cache.set('dashboard_top_routes_v1', chart_top_routes, 300)

    return render(request, 'core/dashboard.html', {
        'active_page': 'dashboard',
        'metrics': {
            'trips_this_month': trips_this_month,
            'trips_active_today': trips_active_today,
            'vehicles_in_use_today': vehicles_in_use_today,
            'tickets_this_month': tickets_this_month,
            'revenue_this_month': revenue_this_month,
            'customers_count': customers_count,
            'vehicles_active': vehicles_active,
            'drivers_available': drivers_available,
            'routes_active': routes_active,
        },
        'upcoming_trips': upcoming_trips,
        'recent_orders': recent_orders,
        'chart_revenue': chart_revenue,
        'chart_today_index': chart_today_index,
        'chart_status': chart_status,
        'chart_top_routes': chart_top_routes,
    })



# рейси


@staff_required
def trips_list(request):
    qs = (
        Trip.objects
        .select_related('route', 'vehicle', 'main_driver__user')
        .annotate(tickets_count=Count(
            'orders__tickets',
            filter=Q(orders__tickets__status__in=['booked', 'paid', 'used']),
        ))
        .order_by('-departure_at')
    )

    status = request.GET.get('status', '')
    route_id = request.GET.get('route', '')
    search = request.GET.get('q', '').strip()
    period = request.GET.get('period', '')

    if status:
        qs = qs.filter(status=status)
    if route_id:
        qs = qs.filter(route_id=route_id)
    if search:
        qs = qs.filter(
            Q(route__name__icontains=search) |
            Q(route__code__icontains=search) |
            Q(vehicle__registration_number__icontains=search)
        )

    # швидкі фільтри по періоду
    now = timezone.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    if period == 'today':
        qs = qs.filter(departure_at__gte=today_start, departure_at__lt=today_start + timedelta(days=1))
    elif period == 'week':
        qs = qs.filter(departure_at__gte=today_start, departure_at__lt=today_start + timedelta(days=7))
    elif period == 'month':
        qs = qs.filter(departure_at__gte=today_start, departure_at__lt=today_start + timedelta(days=30))
    elif period == 'future':
        qs = qs.filter(departure_at__gte=now)
    elif period == 'past':
        qs = qs.filter(departure_at__lt=now)

    paginator = Paginator(qs, PAGE_SIZE)
    page = paginator.get_page(request.GET.get('page'))

    return render(request, 'core/trips_list.html', {
        'active_page': 'trips',
        'page': page,
        'trips': page.object_list,
        'status_choices': Trip.Status.choices,
        'routes': Route.objects.filter(is_active=True).order_by('code'),
        'filters': {'status': status, 'route': route_id, 'q': search, 'period': period},
    })


@staff_required
def trip_detail(request, trip_id):
    trip = get_object_or_404(
        Trip.objects.select_related(
            'route', 'vehicle', 'main_driver__user', 'co_driver__user'
        ),
        pk=trip_id,
    )
    stops = trip.route.stops.all().order_by('order')
    tickets = (
        Ticket.objects
        .filter(order__trip=trip)
        .select_related('order', 'boarding_stop', 'alighting_stop')
        .order_by('passenger_last_name', 'passenger_first_name')
    )
    revenue = sum(
        (t.price for t in tickets if t.status in ('booked', 'paid', 'used')),
        start=Decimal('0'),
    )
    sold = sum(1 for t in tickets if t.status in ('booked', 'paid', 'used'))
    free_seats = max(0, (trip.vehicle.seats_total or 0) - sold)

    return render(request, 'core/trip_detail.html', {
        'active_page': 'trips',
        'trip': trip,
        'stops': stops,
        'tickets': tickets,
        'revenue': revenue,
        'sold': sold,
        'free_seats': free_seats,
    })



# замовлення


@staff_required
def orders_list(request):
    # Обробка bulk-action
    if request.method == 'POST':
        action = request.POST.get('bulk_action')
        ids = request.POST.getlist('selected')
        if action and ids:
            qs = Order.objects.filter(pk__in=ids)
            if action == 'confirm':
                qs.update(status=Order.Status.CONFIRMED)
            elif action == 'mark_paid':
                qs.update(status=Order.Status.PAID, paid_at=timezone.now())
            elif action == 'cancel':
                qs.update(status=Order.Status.CANCELLED)
        from django.http import HttpResponseRedirect
        return HttpResponseRedirect(request.get_full_path())

    qs = (
        Order.objects
        .select_related('trip__route', 'customer', 'created_by')
        .annotate(tickets_total=Count('tickets'))
        .order_by('-created_at')
    )

    status = request.GET.get('status', '')
    search = request.GET.get('q', '').strip()

    if status:
        qs = qs.filter(status=status)
    if search:
        qs = qs.filter(
            Q(order_number__icontains=search) |
            Q(contact_first_name__icontains=search) |
            Q(contact_last_name__icontains=search) |
            Q(contact_phone__icontains=search) |
            Q(customer__name__icontains=search)
        )

    paginator = Paginator(qs, PAGE_SIZE)
    page = paginator.get_page(request.GET.get('page'))

    return render(request, 'core/orders_list.html', {
        'active_page': 'orders',
        'page': page,
        'orders': page.object_list,
        'status_choices': Order.Status.choices,
        'filters': {'status': status, 'q': search},
    })


@staff_required
def order_detail(request, order_id):
    order = get_object_or_404(
        Order.objects.select_related(
            'trip__route', 'trip__vehicle', 'customer', 'created_by'
        ),
        pk=order_id,
    )
    tickets = order.tickets.select_related('boarding_stop', 'alighting_stop').all()
    return render(request, 'core/order_detail.html', {
        'active_page': 'orders',
        'order': order,
        'tickets': tickets,
    })



# клієнти


@staff_required
def customer_detail(request, customer_id):
    """Детальна сторінка клієнта з історією замовлень."""
    customer = get_object_or_404(Customer, pk=customer_id)
    orders = (
        Order.objects
        .filter(customer=customer)
        .select_related('trip__route')
        .annotate(tickets_total=Count('tickets'))
        .order_by('-created_at')[:50]
    )
    stats = Order.objects.filter(customer=customer).aggregate(
        total_orders=Count('id'),
        total_revenue=Sum('total_price'),
    )
    return render(request, 'core/customer_detail.html', {
        'active_page': 'customers',
        'customer': customer,
        'orders': orders,
        'stats': stats,
    })


@staff_required
def customers_list(request):
    qs = (
        Customer.objects
        .annotate(orders_count=Count('orders'))
        .order_by('name')
    )
    country = request.GET.get('country', '')
    search = request.GET.get('q', '').strip()
    if country:
        qs = qs.filter(country=country)
    if search:
        qs = qs.filter(
            Q(name__icontains=search) |
            Q(legal_name__icontains=search) |
            Q(tax_number__icontains=search) |
            Q(contact_person__icontains=search)
        )
    paginator = Paginator(qs, PAGE_SIZE)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'core/customers_list.html', {
        'active_page': 'customers',
        'page': page,
        'customers': page.object_list,
        'countries': Customer.Country.choices,
        'filters': {'country': country, 'q': search},
    })



# автопарк


@staff_required
def vehicles_list(request):
    qs = Vehicle.objects.order_by('vehicle_type', 'registration_number')
    vehicle_type = request.GET.get('type', '')
    is_active = request.GET.get('active', '')
    search = request.GET.get('q', '').strip()
    if vehicle_type:
        qs = qs.filter(vehicle_type=vehicle_type)
    if is_active in ('1', '0'):
        qs = qs.filter(is_active=(is_active == '1'))
    if search:
        qs = qs.filter(
            Q(registration_number__icontains=search) |
            Q(brand__icontains=search) |
            Q(model__icontains=search)
        )
    paginator = Paginator(qs, PAGE_SIZE)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'core/vehicles_list.html', {
        'active_page': 'vehicles',
        'page': page,
        'vehicles': page.object_list,
        'types': Vehicle.VehicleType.choices,
        'filters': {'type': vehicle_type, 'active': is_active, 'q': search},
    })



# водії


@staff_required
def drivers_list(request):
    qs = Driver.objects.select_related('user').order_by(
        'user__last_name', 'user__first_name',
    )
    available = request.GET.get('available', '')
    search = request.GET.get('q', '').strip()
    if available in ('1', '0'):
        qs = qs.filter(is_available=(available == '1'))
    if search:
        qs = qs.filter(
            Q(user__last_name__icontains=search) |
            Q(user__first_name__icontains=search) |
            Q(license_number__icontains=search) |
            Q(license_categories__icontains=search)
        )
    paginator = Paginator(qs, PAGE_SIZE)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'core/drivers_list.html', {
        'active_page': 'drivers',
        'page': page,
        'drivers': page.object_list,
        'filters': {'available': available, 'q': search},
    })



# маршрути



# календар рейсів


STATUS_COLORS = {
    'planned': '#94a3b8',
    'on_sale': '#0ea5e9',
    'in_progress': '#7c3aed',
    'completed': '#10b981',
    'cancelled': '#ef4444',
}


@staff_required
def trips_calendar(request):
    """Сторінка з календарем рейсів."""
    return render(request, 'core/trips_calendar.html', {
        'active_page': 'calendar',
    })


@staff_required
def trips_calendar_feed(request):
    """JSON для FullCalendar з усіма рейсами."""
    start = request.GET.get('start')
    end = request.GET.get('end')
    qs = Trip.objects.select_related('route', 'vehicle').all()
    if start:
        qs = qs.filter(departure_at__gte=start)
    if end:
        qs = qs.filter(departure_at__lt=end)
    events = []
    for t in qs:
        end_at = t.departure_at + timedelta(minutes=t.route.duration_minutes)
        events.append({
            'id': t.id,
            'title': f'{t.route.code} · {t.route.origin_city}-{t.route.destination_city}',
            'start': t.departure_at.isoformat(),
            'end': end_at.isoformat(),
            'url': f'/manage/trips/{t.id}/',
            'backgroundColor': STATUS_COLORS.get(t.status, '#4f46e5'),
            'borderColor': STATUS_COLORS.get(t.status, '#4f46e5'),
        })
    return JsonResponse(events, safe=False)



# експорт в Excel


def _excel_response(filename):
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# символи з яких починається XLSX-формула. якщо у текстовому полі
# зберегти "=cmd|'/c calc'!A1", Excel виконає це при відкритті файла.
# тому усі вхідні рядки що починаються з цих символів префіксуємо
# апострофом - Excel сприймає такий рядок як звичайний текст.
_XLSX_INJECT_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


def _safe(value):
    """захист від xlsx-injection. застосовую до всіх рядкових значень
    які можуть прийти з користувацького вводу (імена, телефони, нотатки)."""
    if isinstance(value, str) and value and value[0] in _XLSX_INJECT_PREFIXES:
        return "'" + value
    return value


@staff_required
def orders_export(request):
    """Експорт усіх замовлень у XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = 'Замовлення'

    headers = [
        '№ замовлення', 'Дата', 'Маршрут', 'Відправлення',
        'Контакт', 'Телефон', 'Клієнт', 'Квитків',
        'Сума', 'Знижка', 'Валюта', 'Статус',
    ]
    ws.append(headers)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4F46E5')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    qs = (
        Order.objects
        .select_related('trip__route', 'customer')
        .annotate(tickets_total=Count('tickets'))
        .order_by('-created_at')
    )
    # iterator(chunk_size=1000) тримає у пам'яті лише 1000 об'єктів за раз
    # замість всіх 23k+ замовлень. Без цього експорт на великих обсягах
    # роздуває пам'ять воркера до 500+ МБ.
    for o in qs.iterator(chunk_size=1000):
        ws.append([
            _safe(o.order_number),
            o.created_at.replace(tzinfo=None),
            _safe(o.trip.route.code if o.trip else ''),
            o.trip.departure_at.replace(tzinfo=None) if o.trip else None,
            _safe(f'{o.contact_last_name} {o.contact_first_name}'),
            _safe(o.contact_phone),
            _safe(o.customer.name if o.customer else ''),
            o.tickets_total,
            float(o.total_price),
            float(o.discount_amount),
            o.currency,
            o.get_status_display(),
        ])

    column_widths = [16, 18, 12, 18, 22, 16, 22, 9, 12, 12, 8, 16]
    for i, w in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    response = _excel_response('orders.xlsx')
    wb.save(response)
    return response


@staff_required
def trips_export(request):
    """Експорт рейсів у XLSX."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = 'Рейси'

    headers = [
        'Маршрут', 'Назва', 'Відправлення', 'Транспорт',
        'Водій', 'Місць', 'Продано', 'Ціна', 'Валюта', 'Статус',
    ]
    ws.append(headers)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4F46E5')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    qs = (
        Trip.objects
        .select_related('route', 'vehicle', 'main_driver__user')
        .annotate(sold=Count(
            'orders__tickets',
            filter=Q(orders__tickets__status__in=['booked', 'paid', 'used']),
        ))
        .order_by('-departure_at')
    )
    # iterator(chunk_size=500) для економії пам'яті. рейсів ~1500, але
    # запас на майбутнє коли база виросте.
    for t in qs.iterator(chunk_size=500):
        driver_name = ''
        if t.main_driver_id:
            driver_name = f'{t.main_driver.user.last_name} {t.main_driver.user.first_name}'
        ws.append([
            _safe(t.route.code),
            _safe(t.route.name),
            t.departure_at.replace(tzinfo=None),
            _safe(f'{t.vehicle.brand} {t.vehicle.model} ({t.vehicle.registration_number})'),
            _safe(driver_name),
            t.vehicle.seats_total or 0,
            t.sold,
            float(t.base_price),
            t.currency,
            t.get_status_display(),
        ])

    column_widths = [12, 30, 18, 32, 28, 8, 10, 10, 8, 16]
    for i, w in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    response = _excel_response('trips.xlsx')
    wb.save(response)
    return response



# сповіщення (для bell-індикатора у топбарі)


@staff_required
def audit_log(request):
    """Сторінка журналу змін системи."""
    qs = AuditLog.objects.select_related('user').order_by('-timestamp')
    action = request.GET.get('action', '')
    model = request.GET.get('model', '')
    search = request.GET.get('q', '').strip()
    if action:
        qs = qs.filter(action=action)
    if model:
        qs = qs.filter(model_name=model)
    if search:
        qs = qs.filter(
            Q(object_repr__icontains=search) |
            Q(model_name__icontains=search)
        )
    paginator = Paginator(qs, 30)
    page = paginator.get_page(request.GET.get('page'))
    models_list = AuditLog.objects.values_list('model_name', flat=True).distinct().order_by('model_name')
    return render(request, 'core/audit_log.html', {
        'active_page': 'audit',
        'page': page,
        'logs': page.object_list,
        'action_choices': AuditLog.Action.choices,
        'models': models_list,
        'filters': {'action': action, 'model': model, 'q': search},
    })


@staff_required
@ratelimit(key='user', rate='120/m', block=True)
def notifications_data(request):
    """JSON з лічильником і списком останніх сповіщень для топбара.
    Показуємо тільки нові замовлення за останні 24 години (інакше лічильник
    накопичує тисячі історичних pending-замовлень).

    rate-limit: 120 запитів/хв на користувача (топбар поллит раз на 30 сек,
    тобто реально 2/хв - 120 з запасом)."""
    since = timezone.now() - timedelta(hours=24)
    recent_pending = Order.objects.filter(
        status=Order.Status.PENDING,
        created_at__gte=since,
    )
    items = []
    for o in recent_pending.select_related('trip__route').order_by('-created_at')[:10]:
        items.append({
            'type': 'order',
            'title': f'Нове замовлення {o.order_number}',
            'subtitle': f'{o.trip.route.code} · {o.contact_last_name} {o.contact_first_name}',
            'created_at': o.created_at.isoformat(),
            'url': f'/manage/orders/{o.id}/',
        })
    return JsonResponse({
        'count': recent_pending.count(),
        'items': items,
    })


PERIOD_OPTIONS = [
    ('7', 'Останні 7 днів', 7),
    ('30', 'Останні 30 днів', 30),
    ('90', 'Останні 90 днів', 90),
    ('365', 'Останні 12 місяців', 365),
]


@staff_required
def reports(request):
    """Сторінка зведених звітів з інтерактивними фільтрами періоду та маршруту."""
    from django.db.models.functions import ExtractIsoWeekDay

    now = timezone.now()

    period_key = request.GET.get('period', '30')
    period_days = next((d for k, _, d in PERIOD_OPTIONS if k == period_key), 30)
    period_start = now - timedelta(days=period_days)

    route_filter = request.GET.get('route', '')

    paid_via_trips = Q(trips__orders__tickets__status__in=['booked', 'paid', 'used'])

    orders_qs = Order.objects.filter(created_at__gte=period_start)
    if route_filter:
        orders_qs = orders_qs.filter(trip__route_id=route_filter)

    revenue_total = orders_qs.filter(
        status__in=['paid', 'in_progress', 'completed']
    ).aggregate(total=Sum('total_price'))['total'] or 0

    tickets_in_period = Ticket.objects.filter(
        created_at__gte=period_start,
        status__in=['booked', 'paid', 'used'],
    )
    if route_filter:
        tickets_in_period = tickets_in_period.filter(order__trip__route_id=route_filter)

    recent_stats = {
        'orders_count': orders_qs.count(),
        'tickets_count': tickets_in_period.count(),
        'revenue': revenue_total,
        'avg_check': 0,
    }
    if recent_stats['orders_count'] > 0:
        recent_stats['avg_check'] = float(recent_stats['revenue']) / recent_stats['orders_count']

    top_drivers = (
        Driver.objects
        .select_related('user')
        .annotate(trips_count=Count(
            'trips_as_main',
            filter=Q(trips_as_main__departure_at__gte=period_start),
            distinct=True,
        ))
        .filter(trips_count__gt=0)
        .order_by('-trips_count')[:10]
    )

    top_vehicles = (
        Vehicle.objects
        .annotate(
            trips_count=Count(
                'trips',
                filter=Q(trips__departure_at__gte=period_start),
                distinct=True,
            ),
            revenue=Sum(
                'trips__orders__tickets__price',
                filter=paid_via_trips & Q(trips__departure_at__gte=period_start),
            ),
        )
        .filter(trips_count__gt=0)
        .order_by('-revenue')[:10]
    )

    top_promos = PromoCode.objects.order_by('-times_used')[:10]

    routes_qs = (
        Route.objects
        .annotate(
            trips_count=Count(
                'trips',
                filter=Q(trips__departure_at__gte=period_start),
                distinct=True,
            ),
            revenue=Sum(
                'trips__orders__tickets__price',
                filter=paid_via_trips & Q(trips__departure_at__gte=period_start),
            ),
            tickets_count=Count(
                'trips__orders__tickets',
                filter=paid_via_trips & Q(trips__departure_at__gte=period_start),
                distinct=True,
            ),
        )
        .filter(trips_count__gt=0)
        .order_by('-revenue')
    )
    if route_filter:
        routes_qs = routes_qs.filter(id=route_filter)
    top_routes = routes_qs[:10]

    weekday_qs = (
        orders_qs
        .annotate(wd=ExtractIsoWeekDay('created_at'))
        .values('wd')
        .annotate(count=Count('id'), revenue=Sum('total_price'))
        .order_by('wd')
    )
    weekday_names = ['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Нд']
    weekday_data = {row['wd']: row for row in weekday_qs}
    chart_weekday = []
    for i in range(1, 8):
        row = weekday_data.get(i, {})
        chart_weekday.append({
            'label': weekday_names[i - 1],
            'count': row.get('count', 0) or 0,
            'revenue': float(row.get('revenue', 0) or 0),
        })

    order_status_qs = orders_qs.values('status').annotate(c=Count('id')).order_by()
    order_status_labels = dict(Order.Status.choices)
    chart_order_status = [
        {'label': order_status_labels.get(r['status'], r['status']), 'value': r['c']}
        for r in order_status_qs
    ]

    return render(request, 'core/reports.html', {
        'active_page': 'reports',
        'top_drivers': top_drivers,
        'top_vehicles': top_vehicles,
        'top_promos': top_promos,
        'top_routes': top_routes,
        'recent_stats': recent_stats,
        'chart_weekday': chart_weekday,
        'chart_order_status': chart_order_status,
        'periods': [(k, label) for k, label, _ in PERIOD_OPTIONS],
        'current_period': period_key,
        'all_routes': Route.objects.filter(is_active=True).order_by('code'),
        'current_route': route_filter,
    })


@staff_required
def routes_list(request):
    qs = (
        Route.objects
        .annotate(
            stops_count=Count('stops', distinct=True),
            trips_count=Count('trips', distinct=True),
        )
        .order_by('code')
    )
    is_active = request.GET.get('active', '')
    search = request.GET.get('q', '').strip()
    if is_active in ('1', '0'):
        qs = qs.filter(is_active=(is_active == '1'))
    if search:
        qs = qs.filter(
            Q(code__icontains=search) |
            Q(name__icontains=search) |
            Q(origin_city__icontains=search) |
            Q(destination_city__icontains=search)
        )
    paginator = Paginator(qs, PAGE_SIZE)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'core/routes_list.html', {
        'active_page': 'routes',
        'page': page,
        'routes': page.object_list,
        'filters': {'active': is_active, 'q': search},
    })
